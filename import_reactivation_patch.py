"""Importacao simples de leads antigos para reativacao comercial.

Aceita CSV ou Excel (.xlsx) com apenas Nome e Telefone/WhatsApp.
Cada contato valido vira Novo Lead, evita duplicados por telefone normalizado
e usa a distribuicao automatica do CRM.
"""
import csv
import io

from flask import flash, redirect, render_template, request, url_for

import patched_app as p

crm = p.crm
app = p.app


def _clean_phone(value):
    raw = str(value or "").strip()
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]
    digits = "".join(ch for ch in raw if ch.isdigit())
    if not digits:
        return ""
    if not digits.startswith("55"):
        digits = "55" + digits
    return digits


def _existing_phones():
    result = set()
    for lead in crm.Lead.query.all():
        phone = _clean_phone(lead.phone)
        if phone:
            result.add(phone)
    return result


def _rows_from_csv(file_storage):
    text = file_storage.stream.read().decode("utf-8-sig", errors="ignore")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    except Exception:
        reader = csv.DictReader(io.StringIO(text), delimiter=";")
    return list(reader)


def _rows_from_xlsx(file_storage):
    from openpyxl import load_workbook

    wb = load_workbook(file_storage.stream, read_only=True, data_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    try:
        headers = [str(v or "").strip().lower() for v in next(values)]
    except StopIteration:
        return []

    rows = []
    for values_row in values:
        rows.append({headers[i]: values_row[i] if i < len(values_row) else "" for i in range(len(headers))})
    return rows


def _pick(row, *keys):
    normalized = {str(k).strip().lower(): v for k, v in row.items()}
    for key in keys:
        value = normalized.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def import_leads_reactivation():
    if not p.session.get("user_id"):
        return redirect(url_for("login"))
    if p.session.get("role") not in ["admin", "manager"]:
        flash("Acesso restrito.", "danger")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        uploaded = request.files.get("file")
        if not uploaded or not uploaded.filename:
            flash("Selecione uma planilha CSV ou Excel.", "danger")
            return redirect(url_for("import_leads"))

        filename = uploaded.filename.lower()
        try:
            if filename.endswith(".xlsx"):
                rows = _rows_from_xlsx(uploaded)
            elif filename.endswith(".csv"):
                rows = _rows_from_csv(uploaded)
            else:
                flash("Formato não aceito. Use CSV ou Excel (.xlsx).", "danger")
                return redirect(url_for("import_leads"))
        except Exception as exc:
            app.logger.exception("Falha ao ler arquivo de importacao")
            flash(f"Não foi possível ler a planilha: {exc}", "danger")
            return redirect(url_for("import_leads"))

        existing = _existing_phones()
        created = 0
        duplicates = 0
        invalid = 0

        for row in rows:
            name = _pick(row, "nome", "name")
            phone = _clean_phone(_pick(row, "telefone", "phone", "whatsapp", "celular"))

            if not name or not phone or len(phone) < 12:
                invalid += 1
                continue
            if phone in existing:
                duplicates += 1
                continue

            lead = crm.Lead(
                name=name[:160],
                phone=phone,
                source="Importação - Reativação",
                investment=0,
                timeframe="Sem prazo",
                meeting_interest="Não informado",
                stage="Novo Lead",
                notes="[Q_REACTIVATION]=1",
            )
            crm.db.session.add(lead)
            crm.db.session.flush()
            crm.assign_round_robin(lead)
            lead.score = 0
            lead.temperature = "Frio"
            lead.stage = "Novo Lead"
            crm.db.session.add(crm.AutomationLog(
                lead_id=lead.id,
                action="Lead importado para reativação",
                detail="Criado como Novo Lead a partir de planilha Nome + Telefone.",
            ))
            existing.add(phone)
            created += 1

        crm.db.session.commit()
        flash(
            f"Importação concluída: {created} novos leads, {duplicates} duplicados ignorados e {invalid} linhas inválidas.",
            "success",
        )
        return redirect(url_for("leads"))

    return render_template("import.html")


# Substitui apenas a função do endpoint existente /import, sem criar rota duplicada.
app.view_functions["import_leads"] = import_leads_reactivation
